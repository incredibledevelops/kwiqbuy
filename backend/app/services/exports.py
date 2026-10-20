import csv
import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font

from app.crud import analytics as analytics_crud


def _fmt_dt(v) -> str:
    if isinstance(v, datetime):
        return v.isoformat()
    return str(v) if v is not None else ""


async def build_summary_csv(days: int) -> bytes:
    start, end = analytics_crud.default_range(days=days)
    summary = await analytics_crud.summary(start, end)
    daily = await analytics_crud.daily_series(start, end)
    vendors = await analytics_crud.top_vendors(start, end, limit=50)
    zones = await analytics_crud.top_zones(start, end, limit=50)

    buf = io.StringIO()
    w = csv.writer(buf)

    w.writerow(["KwiqBuy Analytics Export"])
    w.writerow(["Generated", datetime.now(timezone.utc).isoformat()])
    w.writerow(["Range start", _fmt_dt(start)])
    w.writerow(["Range end", _fmt_dt(end)])
    w.writerow(["Days", days])
    w.writerow([])

    w.writerow(["Summary"])
    for k, v in summary.items():
        w.writerow([k, _fmt_dt(v) if isinstance(v, datetime) else v])
    w.writerow([])

    w.writerow(["Daily"])
    w.writerow(["date", "revenue", "orders"])
    for row in daily:
        w.writerow([row["date"], row["revenue"], row["orders"]])
    w.writerow([])

    w.writerow(["Top vendors"])
    w.writerow(["vendor_id", "vendor_name", "orders", "revenue", "platform_fees"])
    for v in vendors:
        w.writerow([v["vendor_id"], v["vendor_name"], v["orders"], v["revenue"], v["platform_fees"]])
    w.writerow([])

    w.writerow(["Top zones"])
    w.writerow(["zone", "orders", "revenue"])
    for z in zones:
        w.writerow([z["zone"], z["orders"], z["revenue"]])

    return buf.getvalue().encode("utf-8")


async def build_summary_xlsx(days: int) -> bytes:
    start, end = analytics_crud.default_range(days=days)
    summary = await analytics_crud.summary(start, end)
    daily = await analytics_crud.daily_series(start, end)
    vendors = await analytics_crud.top_vendors(start, end, limit=50)
    zones = await analytics_crud.top_zones(start, end, limit=50)

    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Summary"
    ws.append(["KwiqBuy Analytics Export"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Generated", datetime.now(timezone.utc).isoformat()])
    ws.append(["Days", days])
    ws.append(["Start", _fmt_dt(start)])
    ws.append(["End", _fmt_dt(end)])
    ws.append([])
    ws.append(["Key", "Value"])
    ws["A8"].font = bold
    ws["B8"].font = bold
    for k, v in summary.items():
        ws.append([k, _fmt_dt(v) if isinstance(v, datetime) else v])

    ws2 = wb.create_sheet("Daily")
    ws2.append(["date", "revenue", "orders"])
    for c in ws2[1]:
        c.font = bold
    for row in daily:
        ws2.append([row["date"], row["revenue"], row["orders"]])

    ws3 = wb.create_sheet("Top Vendors")
    ws3.append(["vendor_id", "vendor_name", "orders", "revenue", "platform_fees"])
    for c in ws3[1]:
        c.font = bold
    for v in vendors:
        ws3.append([v["vendor_id"], v["vendor_name"], v["orders"], v["revenue"], v["platform_fees"]])

    ws4 = wb.create_sheet("Top Zones")
    ws4.append(["zone", "orders", "revenue"])
    for c in ws4[1]:
        c.font = bold
    for z in zones:
        ws4.append([z["zone"], z["orders"], z["revenue"]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
